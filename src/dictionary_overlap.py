from pathlib import Path
import csv

from openpyxl import load_workbook


LM_PATH = Path(
    "data/dictionary/Loughran-McDonald_MasterDictionary_1993-2025.xlsx"
)

H_PATH = Path(
    "data/dictionary/HIV-4.csv"
)


print("=" * 80)
print("DICTIONARY WORD OVERLAP")
print("=" * 80)

# ------------------------------------------------------------
# LM dictionary
# ------------------------------------------------------------

print("Reading LM dictionary...")

wb = load_workbook(
    LM_PATH,
    read_only=True,
    data_only=True,
)

ws = wb.active

rows = ws.iter_rows(values_only=True)

header = next(rows)

header = [
    str(x).strip()
    if x is not None
    else ""
    for x in header
]

word_i = header.index("Word")
negative_i = header.index("Negative")
positive_i = header.index("Positive")

lm_positive = set()
lm_negative = set()

for row in rows:

    if not row:
        continue

    word = row[word_i]

    if word is None:
        continue

    word = str(word).strip().upper()

    if not word:
        continue

    neg = row[negative_i]
    pos = row[positive_i]

    try:
        if float(pos or 0) > 0:
            lm_positive.add(word)
    except Exception:
        pass

    try:
        if float(neg or 0) > 0:
            lm_negative.add(word)
    except Exception:
        pass

wb.close()

# ------------------------------------------------------------
# Harvard dictionary
# ------------------------------------------------------------

print("Reading Harvard-IV-4 dictionary...")

h_positive = set()
h_negative = set()

with H_PATH.open(
    encoding="utf-8",
    newline="",
) as f:

    reader = csv.DictReader(f)

    for row in reader:

        word = str(
            row.get("Entry", "")
        ).strip().upper()

        if not word:
            continue

        if row.get("Positiv"):
            h_positive.add(word)

        if row.get("Negativ"):
            h_negative.add(word)

# ------------------------------------------------------------
# Results
# ------------------------------------------------------------

print()
print("LM Positive words      :", len(lm_positive))
print("Harvard Positive words :", len(h_positive))

print("LM Negative words      :", len(lm_negative))
print("Harvard Negative words :", len(h_negative))

print()
print("Positive ∩ Positive:")
print(len(lm_positive & h_positive))

print()
print("Negative ∩ Negative:")
print(len(lm_negative & h_negative))

print()
print("Harvard Negative ∩ LM Positive:")
print(len(h_negative & lm_positive))

print()
print("Harvard Positive ∩ LM Negative:")
print(len(h_positive & lm_negative))

print()
print("=" * 80)
print("OVERLAP COMPLETE")
print("=" * 80)
